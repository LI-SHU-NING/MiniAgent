"""文档解析工具：Word(.docx) / PDF / Excel(.xlsx)

统一入口 read_document，按扩展名分发；与文件工具共用工作目录沙箱。
"""
from __future__ import annotations

from pathlib import Path

from .file_tools import _resolve
from .registry import tool


@tool(
    description=(
        "读取文档文件内容，支持 Word(.docx)、PDF、Excel(.xlsx)。"
        "PDF 按页返回文本，Word 含段落与表格，Excel 按工作表返回行。"
        "二进制文档（如网页上传的附件）请先用本工具读取再做分析。"
    )
)
def read_document(path: str, max_chars: int = 10000) -> str:
    """解析文档并返回纯文本，超长自动截断。path 相对工作目录。"""
    p = _resolve(path)
    if not p.exists():
        return f"文件不存在: {p}"
    ext = p.suffix.lower()
    try:
        if ext == ".docx":
            text = _read_docx(p)
        elif ext == ".pdf":
            text = _read_pdf(p)
        elif ext in (".xlsx", ".xlsm"):
            text = _read_xlsx(p)
        elif ext == ".doc":
            return "暂不支持旧版 .doc 格式，请先用 Word 另存为 .docx 再读取。"
        elif ext in (".pptx", ".ppt"):
            return "暂不支持 PPT 格式解析，可将内容导出为 PDF 后读取。"
        else:
            return f"不支持的文档格式: {ext}（支持 .docx / .pdf / .xlsx）"
    except Exception as e:  # noqa: BLE001
        return f"解析失败: {type(e).__name__}: {e}（文件可能损坏或加密）"

    text = text.strip()
    if not text:
        return "文档内容为空（可能是扫描件/图片型 PDF，暂无 OCR 能力）。"
    if len(text) > max_chars:
        text = text[:max_chars] + f"\n...[截断，全文共 {len(text)} 字符。可加大 max_chars 或分次读取]"
    return text


def _read_docx(p: Path) -> str:
    from docx import Document

    doc = Document(str(p))
    parts: list[str] = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text.strip())
    for ti, table in enumerate(doc.tables, 1):
        parts.append(f"[表格{ti}]")
        for row in table.rows:
            cells = [c.text.strip().replace("\n", " ") for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def _read_pdf(p: Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(p))
    parts = []
    for i, page in enumerate(reader.pages, 1):
        txt = (page.extract_text() or "").strip()
        if txt:
            parts.append(f"--- 第 {i} 页 ---\n{txt}")
    if not parts:
        return ""  # 无任何可提取文本 -> 上层提示"内容为空（可能是扫描件）"
    return f"[共 {len(reader.pages)} 页]\n" + "\n".join(parts)


def _read_xlsx(p: Path) -> str:
    from openpyxl import load_workbook

    wb = load_workbook(str(p), read_only=True, data_only=True)
    parts: list[str] = []
    for ws in wb.worksheets:
        parts.append(f"--- 工作表: {ws.title} ---")
        for row in ws.iter_rows(max_row=300, values_only=True):
            line = " | ".join("" if v is None else str(v) for v in row).rstrip(" |")
            if line.strip():
                parts.append(line)
    wb.close()
    return "\n".join(parts)
