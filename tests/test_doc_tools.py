"""文档解析工具测试（离线，无需模型服务）"""
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import agent.tools.file_tools as ft  # noqa: E402
from agent.tools.doc_tools import read_document  # noqa: E402


def test_read_docx(tmp_path):
    ft.set_workspace(str(tmp_path))
    from docx import Document

    doc = Document()
    doc.add_paragraph("这是一个测试段落，包含关键信息 alpha123。")
    doc.add_paragraph("第二段内容。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "姓名"
    table.cell(0, 1).text = "分数"
    table.cell(1, 0).text = "小明"
    table.cell(1, 1).text = "95"
    p = tmp_path / "test.docx"
    doc.save(str(p))

    out = read_document(str(p))
    assert "alpha123" in out
    assert "小明 | 95" in out


def test_read_xlsx(tmp_path):
    ft.set_workspace(str(tmp_path))
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "成绩表"
    ws.append(["姓名", "科目", "分数"])
    ws.append(["小红", "数学", 98])
    p = tmp_path / "test.xlsx"
    wb.save(str(p))

    out = read_document(str(p))
    assert "--- 工作表: 成绩表 ---" in out
    assert "小红 | 数学 | 98" in out


def test_read_pdf_dispatch(tmp_path):
    """PDF 分发逻辑：空 PDF 返回内容为空提示；不支持的格式给出说明"""
    ft.set_workspace(str(tmp_path))
    from pypdf import PdfWriter

    p = tmp_path / "empty.pdf"
    w = PdfWriter()
    w.add_blank_page(width=200, height=200)
    with open(p, "wb") as f:
        w.write(f)
    out = read_document(str(p))
    assert "内容为空" in out or "扫描件" in out

    out2 = read_document(str(tmp_path / "no.doc"))
    assert ".doc" in out2

    (tmp_path / "x.zip").write_bytes(b"PK")
    out3 = read_document(str(tmp_path / "x.zip"))
    assert "不支持" in out3


def test_read_document_missing(tmp_path):
    ft.set_workspace(str(tmp_path))
    out = read_document(str(tmp_path / "ghost.pdf"))
    assert "文件不存在" in out
